# ---------------------------------------------------------------------------
# File   : make_figures_pptx.py
# Author : yoobeom.kim@samsung.com
# Purpose: Editable PowerPoint version of the patent drawings (fig1..fig12).
#          Every element is a native shape, connector, table or chart so the
#          drawings can be revised in PowerPoint.  Same geometry and data as
#          make_figures.py; data figures read example/out results.
#            python3 make_figures_pptx.py [output.pptx]
#          After editing, export the slides as PNG from PowerPoint to replace
#          figures/figN.png used by make_docx.js.
# ---------------------------------------------------------------------------
import csv, os, sys, math
from collections import defaultdict
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_PATTERN_TYPE
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import (XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION,
                             XL_TICK_MARK, XL_MARKER_STYLE)
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.normpath(os.path.join(HERE, "..", "..", "example", "out"))
OUT = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.join(HERE, "patent_figures.pptx")

FONT = "Malgun Gothic"
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x7A, 0x7A, 0x7A)
LIGHT = RGBColor(0xD9, 0xD9, 0xD9)
SLIDE_W, SLIDE_H = 10.0, 7.5
AREA = (0.5, 0.9, 9.0, 6.3)            # drawing area: left, top, width, height [in]
SG = {("N", "vth"): 0.015, ("P", "vth"): 0.015, ("N", "u0"): 0.03, ("P", "u0"): 0.03}
DESC = {
    1: "도 1은 분리·결합 모델링 구조와 데이터 흐름을 나타내는 도면이다.",
    2: "도 2는 트랜지스터별 개별 섭동 시뮬레이션과 그룹 동시 섭동 시뮬레이션의 개념도이다.",
    3: "도 3은 LVF 특성화 데이터로부터 트랜지스터별 기여도 표면을 생성하는 과정을 나타내는 도면이다.",
    4: "도 4는 집합 민감도 라이브러리로부터 소자 파라미터별 집합 민감도 표면을 추출하는 과정을 나타내는 도면이다.",
    5: "도 5는 트랜지스터별 기여도와 집합 민감도의 결합에 의한 트랜지스터별 민감도 복원과 정합성 조건을 나타내는 도면이다.",
    6: "도 6은 셀 인스턴스별 국부 레이아웃 효과 변동량을 타이밍 아크 보정으로 변환하여 정적 타이밍 분석에 반영하는 흐름도이다.",
    7: "도 7은 동일한 트랜지스터별 민감도로부터 국부 변동 통계량(제곱합의 제곱근)과 전역 변동 타이밍 변화량(부호 합)이 산출되는 것을 나타내는 도면이다.",
    8: "도 8은 정합성 검증 단계의 흐름도이다.",
    9: "도 9는 상대 변화율의 보간을 통한 타이밍 라이브러리 테이블 갱신의 개념도이다.",
    10: "도 10은 실시예에서 전역 변동점의 모델 예측값과 직접 시뮬레이션 값을 비교한 그래프이다.",
    11: "도 11은 실시예에서 기여도와 집합 민감도로 복원한 트랜지스터별 민감도의 오차 및 파라미터별 기여도의 비교를 나타내는 그래프이다.",
    12: "도 12는 장치의 블록도이다.",
}


# ---------------------------------------------------------------------------
# text / shape helpers
# ---------------------------------------------------------------------------
def set_run_font(run, size, bold=False, color=BLACK, name=FONT):
    f = run.font
    f.name = name
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    if latin is not None and rPr.find(qn("a:ea")) is None:
        ea = etree.Element(qn("a:ea"))
        ea.set("typeface", name)
        latin.addnext(ea)


def fill_text(tf, lines, size, align=PP_ALIGN.CENTER, bold=False, color=BLACK):
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        set_run_font(r, size, bold, color)


def text_width_in(s, size):
    n = sum(1.0 if ord(ch) > 0x2E80 else 0.58 for ch in s)
    return n * size / 72.0 + 0.12


def style_shape(shp, fill=WHITE, lw=1.0, dash=False, line=BLACK):
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(lw)
    if dash:
        shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shp.shadow.inherit = False


def add_arrowhead(connector):
    ln = connector.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")


def plain_table(graphic_frame):
    tbl = graphic_frame._element.graphic.graphicData.tbl
    tblPr = tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"      # No Style, Table Grid


def cell_borders(cell, w_pt=0.75, color="000000"):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag))
        ln.set("w", str(int(Pt(w_pt))))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        clr = etree.SubElement(sf, qn("a:srgbClr"))
        clr.set("val", color)
        pd = etree.SubElement(ln, qn("a:prstDash"))
        pd.set("val", "solid")


def add_textbox(slide, left, top, width, height, lines, size, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE, bold=False, color=BLACK, wrap=False):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    fill_text(tf, lines, size, align, bold, color)
    return tb


class Canvas:
    """Maps the figure coordinate system of make_figures.py (x right, y up)
    onto the slide drawing area, preserving aspect ratio."""

    def __init__(self, slide, xl, yl, mpl_w, area=AREA, fs_min=8.0):
        self.slide, self.xl, self.yl = slide, xl, yl
        L, T, W, H = area
        self.s = min(W / xl, H / yl)
        self.ox = L + (W - xl * self.s) / 2
        self.oy = T + (H - yl * self.s) / 2
        self.factor = self.s / (mpl_w / xl)
        self.fs_min = fs_min

    def xin(self, x): return self.ox + x * self.s
    def yin(self, y): return self.oy + (self.yl - y) * self.s
    def X(self, x): return Inches(self.xin(x))
    def Y(self, y): return Inches(self.yin(y))
    def L(self, v): return Inches(v * self.s)
    def fs(self, f): return max(self.fs_min, round(f * self.factor * 2) / 2)

    def rect(self, x, y, w, h, text="", fs=8.5, fill=WHITE, lw=1.0, dash=False,
             shape=MSO_SHAPE.RECTANGLE):
        shp = self.slide.shapes.add_shape(shape, self.X(x), self.Y(y + h), self.L(w), self.L(h))
        style_shape(shp, fill, lw, dash)
        tf = shp.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        fill_text(tf, text.split("\n") if text else [""], self.fs(fs))
        return shp

    def diamond(self, cx, cy, w, h, text, fs=8.5):
        return self.rect(cx - w / 2, cy - h / 2, w, h, text, fs, shape=MSO_SHAPE.DIAMOND)

    def text(self, x, y, s, fs=8, ha="center", va="center", bold=False, color=BLACK):
        size = self.fs(fs)
        lines = s.split("\n")
        w = max(text_width_in(l, size) for l in lines)
        h = len(lines) * size / 72.0 * 1.3 + 0.04
        xi, yi = self.xin(x), self.yin(y)
        left = {"center": xi - w / 2, "left": xi, "right": xi - w}[ha]
        top = {"center": yi - h / 2, "top": yi, "bottom": yi - h}[va]
        align = {"center": PP_ALIGN.CENTER, "left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}[ha]
        anchor = {"center": MSO_ANCHOR.MIDDLE, "top": MSO_ANCHOR.TOP, "bottom": MSO_ANCHOR.BOTTOM}[va]
        return add_textbox(self.slide, left, top, w, h, lines, size, align, anchor, bold, color)

    def line(self, x0, y0, x1, y1, lw=1.0, dash=False, color=BLACK, arrow=False):
        c = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, self.X(x0), self.Y(y0), self.X(x1), self.Y(y1))
        c.line.color.rgb = color
        c.line.width = Pt(lw)
        if dash:
            c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if arrow:
            add_arrowhead(c)
        return c

    def arrow(self, x0, y0, x1, y1, text="", tx=None, ty=None, fs=8, dash=False):
        c = self.line(x0, y0, x1, y1, dash=dash, arrow=True)
        if text:
            self.text(tx if tx is not None else (x0 + x1) / 2 + 0.15,
                      ty if ty is not None else (y0 + y1) / 2, text, fs, ha="left", va="center")
        return c

    def table(self, x0, y_top, cws, rh, hdr, rows, fs=7.5, hdr_fill=LIGHT):
        """native table; (x0, y_top) = top-left corner in figure units, rows go downward"""
        allrows = ([hdr] if hdr else []) + list(rows)
        nrows, ncols = len(allrows), len(cws)
        gf = self.slide.shapes.add_table(nrows, ncols, self.X(x0), self.Y(y_top), self.L(sum(cws)), self.L(rh * nrows))
        plain_table(gf)
        tbl = gf.table
        for c, w in enumerate(cws):
            tbl.columns[c].width = self.L(w)
        for r in range(nrows):
            tbl.rows[r].height = self.L(rh)
        for r, vals in enumerate(allrows):
            for c, v in enumerate(vals):
                cell = tbl.cell(r, c)
                cell_borders(cell)
                cell.fill.solid()
                cell.fill.fore_color.rgb = hdr_fill if (hdr and r == 0) else WHITE
                cell.margin_left = cell.margin_right = Inches(0.03)
                cell.margin_top = cell.margin_bottom = Inches(0.01)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                fill_text(cell.text_frame, [v], self.fs(fs))
        return gf


def new_slide(prs, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide, 0.5, 0.3, 9.0, 0.45, [f"[도 {n}]"], 14)
    slide.notes_slide.notes_text_frame.text = DESC[n]
    return slide


# ---------------------------------------------------------------------------
# chart helpers
# ---------------------------------------------------------------------------
def chart_common(chart, size=9):
    chart.font.size = Pt(size)
    chart.font.name = FONT
    cs = chart._chartSpace
    if cs.find(qn("c:spPr")) is None:
        spPr = etree.SubElement(cs, qn("c:spPr"))
        etree.SubElement(spPr, qn("a:noFill"))
        ln = etree.SubElement(spPr, qn("a:ln"))
        etree.SubElement(ln, qn("a:noFill"))
        cs.find(qn("c:chart")).addnext(spPr)


def chart_title(chart, text, size=10):
    chart.has_title = True
    tf = chart.chart_title.text_frame
    tf.text = text
    set_run_font(tf.paragraphs[0].runs[0], size, bold=False)


def chart_legend(chart, size=8.5):
    chart.has_legend = True
    lg = chart.legend
    lg.position = XL_LEGEND_POSITION.TOP
    lg.include_in_layout = False
    lg.font.size = Pt(size)
    lg.font.name = FONT


def bar_style(obj, hatch=False, lw=0.75):
    fmt = obj.format
    if hatch:
        fmt.fill.patterned()
        fmt.fill.pattern = MSO_PATTERN_TYPE.WIDE_UPWARD_DIAGONAL
        fmt.fill.fore_color.rgb = BLACK
        fmt.fill.back_color.rgb = WHITE
    else:
        fmt.fill.solid()
        fmt.fill.fore_color.rgb = WHITE
    fmt.line.color.rgb = BLACK
    fmt.line.width = Pt(lw)


def axis_style(ax, size=8.5, title=None, tsize=9, fmt=None, minimum=None, maximum=None,
               unit=None, tick=XL_TICK_MARK.OUTSIDE, rot=None):
    ax.tick_labels.font.size = Pt(size)
    ax.tick_labels.font.name = FONT
    ax.format.line.color.rgb = BLACK
    ax.format.line.width = Pt(0.75)
    ax.has_major_gridlines = False
    ax.has_minor_gridlines = False
    ax.major_tick_mark = tick
    ax.minor_tick_mark = XL_TICK_MARK.NONE
    if fmt:
        ax.tick_labels.number_format = fmt
        ax.tick_labels.number_format_is_linked = False
    if minimum is not None:
        ax.minimum_scale = minimum
    if maximum is not None:
        ax.maximum_scale = maximum
    if unit is not None:
        ax.major_unit = unit
    if title:
        ax.has_title = True
        tf = ax.axis_title.text_frame
        tf.text = title
        set_run_font(tf.paragraphs[0].runs[0], tsize, bold=False)
    if rot is not None:
        txPr = ax._element.get_or_add_txPr()
        txPr.find(qn("a:bodyPr")).set("rot", str(int(rot * 60000)))


def nice_max(v):
    for unit in (0.1, 0.2, 0.5, 1, 2, 5, 10, 20):
        if v / unit <= 8:
            return math.ceil(v / unit) * unit, unit
    return v, v / 5


def add_column_chart(slide, x, y, w, h, cats, series, title, ytitle, ymin=0, ymax=None, unit=None,
                     yfmt="0.0", legend=True, rot=None, gap=60, labels=False, hatch_from=1, cat_size=8.5):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals, number_format="0.00")
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    chart_common(ch)
    chart_title(ch, title)
    if legend:
        chart_legend(ch)
    else:
        ch.has_legend = False
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.overlap = 0
    plot.vary_by_categories = False
    for i, ser in enumerate(plot.series):
        bar_style(ser, hatch=(i >= hatch_from))
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = "0.00"
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(9)
        dl.font.name = FONT
    axis_style(ch.category_axis, size=cat_size, tick=XL_TICK_MARK.NONE, rot=rot)
    axis_style(ch.value_axis, title=ytitle, fmt=yfmt, minimum=ymin, maximum=ymax, unit=unit)
    return ch


# ---------------------------------------------------------------------------
# data
# ---------------------------------------------------------------------------
def read_csv(name):
    with open(os.path.join(DATA, name), newline="") as fh:
        return list(csv.DictReader(fh))


def sens_tables():
    S = defaultdict(dict); Gd = {}; V0 = {}
    for r in read_csv("sens.csv"):
        if r["meas"] != "delay" or r["s1"] == "":
            continue
        k = (r["cell"], r["arc"], r["slew_idx"], r["load_idx"], r["group"], r["param"])
        if r["inst"].startswith("GROUP:"):
            Gd[k] = float(r["s1"])
        else:
            S[k][r["inst"]] = float(r["s1"]); V0[k] = float(r["v0"])
    return S, Gd, V0


# ---------------------------------------------------------------------------
# 도 1 : 분리·결합 모델링 구조
# ---------------------------------------------------------------------------
def fig1(prs):
    cv = Canvas(new_slide(prs, 1), 12, 9.1, 6.6)
    cv.rect(0.1, 7.6, 3.7, 1.3, "LVF 특성화 데이터 (110)\n트랜지스터별 섭동 민감도 S1(i,q)\nTR 식별 정보 있음", 7)
    cv.rect(4.15, 7.6, 3.7, 1.3, "집합 민감도 라이브러리 (120)\nsensitivity 그룹 S_agg(p)\nTR 식별 정보 없음", 7)
    cv.rect(8.2, 7.6, 3.7, 1.3, "레이아웃 / LLE 엔진 (130)\n인스턴스·TR별 ΔVth(i), ΔU0(i)\n(또는 전역 변동, 열화)", 7)
    cv.rect(0.1, 5.4, 3.7, 1.2, "S210 기여도 결정\nC(i; s,l) = |S1(i)| / Σ|S1(j)|\nΣ C(i) = 1", 7.5)
    cv.rect(4.15, 5.4, 3.7, 1.2, "S220 집합 민감도 획득\nS_agg(p; s,l) 표면", 7.5)
    cv.rect(8.2, 5.4, 3.7, 1.2, "S240 변동량 획득\nΔp(i), 경계 + 내부 TR", 7.5)
    for x in (1.95, 6.0, 10.05):
        cv.arrow(x, 7.6, x, 6.6)
    cv.rect(2.3, 3.4, 7.4, 1.1, "S230 트랜지스터별 파라미터 민감도 복원\nS_rec(i,p; s,l) = C(i; s,l) · S_agg(p; s,l)", 8)
    cv.line(1.95, 5.4, 1.95, 4.5); cv.arrow(1.95, 4.5, 2.3, 4.5)
    cv.arrow(6.0, 5.4, 6.0, 4.5)
    cv.rect(2.3, 1.7, 7.4, 1.1, "S250 타이밍 변화량 산출\nΔD_arc = Σ_i Σ_p C(i) · S_agg(p) · Δp(i)", 8)
    cv.arrow(6.0, 3.4, 6.0, 2.8)
    cv.line(10.05, 5.4, 10.05, 2.25); cv.arrow(10.05, 2.25, 9.7, 2.25)
    cv.rect(0.3, 0.2, 5.4, 1.0, "S260a STA 보정\nD_corrected = D_base + ΔD_arc (인스턴스별)", 7.5)
    cv.rect(6.3, 0.2, 5.4, 1.0, "S260b 타이밍 라이브러리 생성\nv_new = v_old(1 + r), 셀별 디레이트", 7.5)
    cv.arrow(4.5, 1.7, 3.0, 1.2); cv.arrow(7.5, 1.7, 9.0, 1.2)


# ---------------------------------------------------------------------------
# 도 2 : 개별 섭동과 그룹 동시 섭동의 개념도
# ---------------------------------------------------------------------------
def fig2(prs):
    cv = Canvas(new_slide(prs, 2), 12, 6.8, 6.3)

    def cell(ox, title, marks):
        cv.text(ox + 1.6, 6.55, title, 9)
        cv.line(ox + 0.4, 6.1, ox + 2.8, 6.1); cv.text(ox + 2.9, 6.1, "VDD", 7, ha="left")
        for name, bx in (("M3", ox + 0.5), ("M4", ox + 1.9)):
            cv.rect(bx, 4.9, 0.8, 0.8, name + "\n(P)", 7, fill=LIGHT if name in marks else WHITE)
            cv.line(bx + 0.4, 5.7, bx + 0.4, 6.1); cv.line(bx + 0.4, 4.4, bx + 0.4, 4.9)
        cv.line(ox + 0.9, 4.4, ox + 2.3, 4.4); cv.line(ox + 1.6, 3.7, ox + 1.6, 4.4)
        cv.text(ox + 1.75, 4.05, "ZN", 7, ha="left")
        cv.rect(ox + 1.2, 2.9, 0.8, 0.8, "M1\n(N)", 7, fill=LIGHT if "M1" in marks else WHITE)
        cv.rect(ox + 1.2, 1.6, 0.8, 0.8, "M2\n(N)", 7, fill=LIGHT if "M2" in marks else WHITE)
        cv.line(ox + 1.6, 2.4, ox + 1.6, 2.9); cv.line(ox + 1.6, 1.1, ox + 1.6, 1.6)
        cv.line(ox + 0.9, 1.1, ox + 2.3, 1.1); cv.text(ox + 2.4, 1.1, "VSS", 7, ha="left")
        cv.text(ox + 0.6, 3.3, "A", 7, ha="right"); cv.text(ox + 0.6, 2.0, "B", 7, ha="right")

    cell(0.3, "(a) 개별 섭동 : M1", {"M1"})
    cell(4.3, "(b) 개별 섭동 : M2", {"M2"})
    cell(8.3, "(c) 그룹 동시 섭동 : M1, M2", {"M1", "M2"})
    cv.text(0.3, 0.55, "ΔD(M1) = D+(M1) - D0", 8.5, ha="left", va="bottom")
    cv.text(4.3, 0.55, "ΔD(M2) = D+(M2) - D0", 8.5, ha="left", va="bottom")
    cv.text(8.3, 0.55, "ΔD(N) = D+(N) - D0", 8.5, ha="left", va="bottom")
    cv.text(6, 0.05, "검증 :  ΔD(N)  ≒  ΔD(M1) + ΔD(M2)     (1차,  음영 = q+δ 섭동된 트랜지스터)", 9, va="bottom")


# ---------------------------------------------------------------------------
# 도 3 : LVF 데이터로부터 기여도 표면 생성
# ---------------------------------------------------------------------------
def fig3(prs):
    cv = Canvas(new_slide(prs, 3), 12, 5.6, 6.6)
    cv.rect(0.2, 4.4, 2.3, 1.0, "특성화 덱\narc_data", 7.5)
    cv.rect(2.9, 4.4, 2.3, 1.0, "±δ 섭동 측정 결과\n(D+, D-, D0)", 7.5)
    cv.rect(5.6, 4.4, 2.3, 1.0, "S1(i,q; s,l)\n중앙 차분", 7.5)
    cv.rect(8.3, 4.4, 3.5, 1.0, "기여도 표면 C(i; s,l)\n조회 테이블 / 속성 / 적합 함수", 7.5)
    cv.arrow(2.5, 4.9, 2.9, 4.9); cv.arrow(5.2, 4.9, 5.6, 4.9); cv.arrow(7.9, 4.9, 8.3, 4.9)
    cv.text(6, 3.7, "C(i; s,l) = |S1(i,q; s,l)| / Σ_j |S1(j,q; s,l)|      (부호 유지 시  C(i) = S1(i) / Σ_j S1(j)),   Σ_i C(i) = 1", 8.5)
    hdr = ["slew", "load", "XMP0", "XMP0@2", "XMP0@3"]
    rows = [["S1", "L1", "0.13", "0.13", "0.12"], ["S1", "L2", "0.15", "0.11", "0.14"], ["S2", "L1", "0.10", "0.17", "0.12"]]
    cv.table(2.2, 3.1, [1.5] * 5, 0.5, hdr, rows)
    cv.text(6, 0.35, "기여도는 셀 / 아크 / 천이 방향 / 조건(when) / slew / load 에 의존하는 표면으로 저장된다 (예시 값)", 8, va="bottom")


# ---------------------------------------------------------------------------
# 도 4 : 집합 민감도 라이브러리로부터 S_agg(p) 추출
# ---------------------------------------------------------------------------
def fig4(prs):
    cv = Canvas(new_slide(prs, 4), 12, 5.3, 6.6)
    cv.rect(0.1, 3.3, 3.9, 1.8, "sensitivity 그룹\n• device_param : delta_p_vta ...\n• contributing_devices : all\n• sens_cell_rise / sens_cell_fall", 7)
    cv.rect(4.4, 3.4, 3.4, 1.6, "S_agg(p; s,l) =\n[ΔD(+p) - ΔD(-p)]\n/ [p+ - p-]", 7.5)
    cv.rect(8.2, 3.5, 3.7, 1.4, "표면 또는 대표값\n표면 / mean_abs / RMS / max", 7.5)
    cv.arrow(4.0, 4.2, 4.4, 4.2); cv.arrow(7.8, 4.2, 8.2, 4.2)
    hdr = ["대표값", "용도"]
    rows = [["표면 S(s,l)", "STA 보간, 정확도 우선"], ["mean_abs", "셀 / 아크 대표값, 보고"],
            ["RMS", "큰 민감도 강조"], ["max", "guard-band, 최악 조건"]]
    cv.table(2.6, 3.05, [2.4, 4.4], 0.45, hdr, rows)
    cv.text(6, 0.25, "모든 소자를 동시에 섭동하여 산출되므로 트랜지스터 식별 정보가 없다  →  LVF 기여도가 보완", 8, va="bottom")


# ---------------------------------------------------------------------------
# 도 5 : 결합에 의한 복원과 정합성 조건
# ---------------------------------------------------------------------------
def fig5(prs):
    cv = Canvas(new_slide(prs, 5), 12, 4.7, 6.6)
    cv.rect(0.3, 3.2, 2.6, 1.3, "C(i; s,l)\nTR 가중치, Σ C = 1\n(LVF)", 7.5)
    cv.rect(3.6, 3.2, 2.6, 1.3, "S_agg(p; s,l)\n집합 파라미터 민감도\n(USM / 그룹 섭동)", 7.5)
    cv.rect(6.9, 3.2, 2.6, 1.3, "S_rec(i,p; s,l)\n= C(i) · S_agg(p)\nTR별 파라미터 민감도", 7.5)
    cv.rect(10.0, 3.2, 1.8, 1.3, "STA-ready\n모델", 7.5)
    cv.text(3.25, 3.85, "×", 14)
    cv.arrow(6.2, 3.85, 6.9, 3.85); cv.arrow(9.5, 3.85, 10.0, 3.85)
    cv.text(6, 2.3, "ΔD(i) = S_rec(i,p; s,l) · Δp(i)", 9)
    cv.rect(0.2, 0.4, 11.6, 1.3,
            "정합성 조건 : 모든 TR에 동일한 Δp 가 인가되면  Σ_i C(i) · S_agg(p) · Δp = S_agg(p) · Δp\n"
            "(동시 변동 시 타이밍 변화량 = 개별 TR 변화량의 합,  집합 민감도 = TR별 민감도의 부호 합)", 7.5)


# ---------------------------------------------------------------------------
# 도 6 : 인스턴스별 LLE 변동량 → STA 보정 흐름
# ---------------------------------------------------------------------------
def fig6(prs):
    cv = Canvas(new_slide(prs, 6), 12, 7.2, 6.6)
    cv.rect(0.2, 5.6, 2.6, 1.4, "레이아웃 / LLE 엔진\n인스턴스별\nΔVth(i), ΔU0(i)", 7.5)
    cv.rect(3.3, 5.6, 2.8, 1.4, "조회 / 보간\nC(i; s,l), S_agg(p; s,l)\n(인스턴스의 slew, load)", 7.5)
    cv.rect(6.6, 5.6, 2.6, 1.4, "TR별 ΔD(i)\n파라미터별", 7.5)
    cv.rect(9.7, 5.6, 2.1, 1.4, "아크 보정\nΔD_arc", 7.5)
    cv.arrow(2.8, 6.3, 3.3, 6.3); cv.arrow(6.1, 6.3, 6.6, 6.3); cv.arrow(9.2, 6.3, 9.7, 6.3)
    cv.text(6, 4.7, "ΔD_arc = Σ_i Σ_p C(i) · S_agg(p) · Δp(i,p),      D_corrected = D_base + ΔD_arc", 8.5)
    hdr = ["데이터", "조회 키"]
    rows = [["기여도", "cell / arc / transition / when / slew / load / TR"],
            ["집합 민감도", "cell / arc / transition / parameter / slew / load"],
            ["LLE 변동량", "instance / TR / parameter"],
            ["STA 출력", "보정된 셀 지연 / 천이 시간 (또는 인스턴스별 디레이트)"]]
    cv.table(1.2, 4.3, [2.2, 7.4], 0.5, hdr, rows)
    cv.text(6, 0.9, "경계 트랜지스터와 내부 트랜지스터 모두 반영, 트랜지스터별 LLE 재특성화 불필요", 8, va="bottom")


# ---------------------------------------------------------------------------
# 도 7 : 동일 민감도의 두 가지 결합 (국부 RSS, 전역 부호 합)
# ---------------------------------------------------------------------------
def fig7(prs):
    rows = [r for r in read_csv("contrib.csv")
            if r["cell"] == "NAND2_X1" and r["arc"] == "A_r_ZN_f" and r["meas"] == "delay"
            and r["slew_idx"] == "1" and r["load_idx"] == "1" and r["param"] == "vth"]
    sig = [r for r in read_csv("local_sigma.csv")
           if r["cell"] == "NAND2_X1" and r["arc"] == "A_r_ZN_f" and r["meas"] == "delay"
           and r["slew_idx"] == "1" and r["load_idx"] == "1"][0]
    v0 = float(sig["v0"])
    insts = [r["inst"] for r in rows]
    c_local = [float(r["c_local"]) / v0 * 100 for r in rows]
    c_glob = [float(r["c_global_1sig"]) * 3 / v0 * 100 for r in rows]
    rss = math.sqrt(sum(c * c for c in c_local))
    tot = sum(c_glob)
    cats = [f"{i} ({'N' if i in ('M1', 'M2') else 'P'})" for i in insts]
    slide = new_slide(prs, 7)
    add_textbox(slide, 0.5, 0.85, 9.0, 0.4,
                ["NAND2 A→ZN 하강, 문턱 전압 파라미터, 트랜지스터별 기여 (1σ mismatch / 3σ 전역)"], 11)
    for vals, total, title, tlabel, x in (
            (c_local, rss, "(a) 국부 변동 : 독립 → 제곱합의 제곱근", "RSS", 0.5),
            (c_glob, tot, "(b) 전역 변동 : 공통 이동 → 부호 합", "Σ", 5.2)):
        ymax, unit = nice_max(max(vals + [total]) * 1.3)
        ch = add_column_chart(slide, x, 1.5, 4.3, 4.9, cats + [tlabel], [("기여", vals + [total])],
                              title, "지연 대비 [%]", ymax=ymax, unit=unit, legend=False, labels=True,
                              hatch_from=9, yfmt="0")
        bar_style(ch.plots[0].series[0].points[len(vals)], hatch=True)


# ---------------------------------------------------------------------------
# 도 8 : 정합성 검증 흐름도
# ---------------------------------------------------------------------------
def fig8(prs):
    cv = Canvas(new_slide(prs, 8), 10, 10.6, 6.3)
    x, w, h = 2.0, 6.0, 1.1
    cv.rect(x, 9.4, w, h, "그룹 동시 섭동 시뮬레이션 (±δ)\nG1_direct(g,q), H2_direct(g,q) 산출")
    cv.text(x - 0.15, 9.95, "S141", 9, ha="right")
    cv.rect(x, 7.5, w, 1.3, "차이 산출\nd1 = |G1 - G1_direct|·Δref / D0\nd2 = ½·(H2_direct - H2)·Δref² / D0", 8)
    cv.text(x - 0.15, 8.15, "S142", 9, ha="right")
    cv.arrow(5, 9.4, 5, 8.8)
    cv.diamond(5, 6.2, 3.6, 1.5, "d1 ≤ T1 ?")
    cv.text(x - 0.15, 6.2, "S143", 9, ha="right")
    cv.arrow(5, 7.5, 5, 6.95)
    cv.rect(7.4, 5.65, 2.5, 1.1, "특성화 데이터 오류\n(수렴, 정밀도, 부호)\n재특성화", 7.5)
    cv.arrow(6.8, 6.2, 7.4, 6.2); cv.text(7.1, 6.3, "아니오", 8, va="bottom")
    cv.diamond(5, 3.9, 3.6, 1.5, "|d2| ≤ T2 ?")
    cv.arrow(5, 5.45, 5, 4.65, "예", tx=5.15, ty=5.05)
    cv.rect(7.4, 3.35, 2.5, 1.1, "H2_direct 사용\n또는 직접 특성화 대상", 7.5)
    cv.arrow(6.8, 3.9, 7.4, 3.9); cv.text(7.1, 4.0, "아니오", 8, va="bottom")
    cv.rect(x, 1.5, w, h, "합산 민감도 G1, H2 및 기여도 C(i) 사용 → 타이밍 변화량 산출")
    cv.text(x - 0.15, 2.05, "S144", 9, ha="right")
    cv.arrow(5, 3.15, 5, 2.6, "예", tx=5.15, ty=2.9)
    cv.text(5, 0.7, "T1, T2 : 타이밍 값 대비 비율 임계값 (예 : 0.5 %, 1 %),  Δref : 기준 변동량 (예 : 3σ)", 8)


# ---------------------------------------------------------------------------
# 도 9 : 라이브러리 테이블 갱신 개념도
# ---------------------------------------------------------------------------
def fig9(prs):
    pred = [r for r in read_csv("pred_SSG.csv") if r["cell"] == "INV_X1" and r["arc"] == "A_r_ZN_f" and r["meas"] == "delay"]
    grid = {(int(r["slew_idx"]), int(r["load_idx"])): float(r["dv_rel_pred"]) for r in pred}
    cv = Canvas(new_slide(prs, 9), 12, 5.8, 6.3)
    cv.text(2.2, 5.5, "모델 격자 : 상대 변화율 r(slew, load)", 9)
    slews = ["10", "30", "80"]; loads = ["1", "4", "12"]
    rows = [[f"{grid[(j, k)]:+.3f}" for k in range(3)] for j in range(3)]
    cv.table(0.9, 4.8, [1.0] * 3, 0.9, None, rows)
    for k, l in enumerate(loads):
        cv.text(1.4 + k * 1.0, 4.95, l + " fF", 7)
    for j, s in enumerate(slews):
        cv.text(0.8, 4.35 - j * 0.9, s + " ps", 7, ha="right")
    cv.text(2.4, 1.55, "slew × load, 시뮬레이션 격자", 7.5)
    cv.arrow(4.2, 3.3, 5.5, 3.3)
    cv.text(4.85, 3.55, "이중 선형 보간\n(격자 밖은 경계값)", 7.5, va="bottom")
    cv.text(9.1, 5.5, "라이브러리 테이블 (cell_fall) : v_new = v_old · (1 + r)", 9)
    lib_slew = ["5", "20", "50", "100"]; lib_load = ["0.5", "2", "6", "16"]
    cv.table(6.8, 4.95, [1.15] * 4, 0.75, None, [["v·(1+r)"] * 4 for _ in range(4)], fs=6.5)
    for k, l in enumerate(lib_load):
        cv.text(6.8 + k * 1.15 + 0.575, 5.05, l + " fF", 7)
    for j, s in enumerate(lib_slew):
        cv.text(6.7, 4.2 - j * 0.75 + 0.375, s + " ps", 7, ha="right")
    cv.text(9.1, 0.9, "라이브러리 고유 인덱스 (단위 : time_unit, capacitive_load_unit 로 변환)", 7.5)


# ---------------------------------------------------------------------------
# 도 10 : 전역 변동점 예측값과 직접 시뮬레이션 값의 비교
# ---------------------------------------------------------------------------
def fig10(prs):
    corners = ["NVT_p1s", "NVT_p3s", "NVT_m3s", "PVT_p3s", "PVT_m3s", "NU0_m3s", "PU0_m3s", "SS1", "SSG", "FFG", "SFG", "FSG"]
    labels = ["N vth +1σ", "N vth +3σ", "N vth -3σ", "P vth +3σ", "P vth -3σ", "N u0 -3σ", "P u0 -3σ", "SS1", "SSG", "FFG", "SFG", "FSG"]
    xs_sim, ys_pred, e_lin, e_logq = [], [], [], []
    for c in corners:
        rows = [r for r in read_csv(f"validate_{c}.csv") if r["meas"] == "delay"]
        el, eq = [], []
        for r in rows:
            v0 = float(r["v0"]); vs = float(r["v_sim"])
            xs_sim.append(vs / v0)
            ys_pred.append((v0 + float(r["dv_logquad"])) / v0)
            el.append(abs(float(r["err_abs_pct_lin"])))
            eq.append(abs(float(r["err_abs_pct_logquad"])))
        e_lin.append(sum(el) / len(el)); e_logq.append(sum(eq) / len(eq))
    slide = new_slide(prs, 10)
    # (a) scatter, prediction vs direct simulation
    # points first, y = x reference line second (drawn on top); black filled
    # markers: some viewers do not draw white-filled markers
    cd = XyChartData()
    s0 = cd.add_series("지연 포인트")
    for xv, yv in zip(xs_sim, ys_pred):
        s0.add_data_point(round(xv, 5), round(yv, 5))
    s1 = cd.add_series("y = x")
    s1.add_data_point(0.8, 0.8); s1.add_data_point(1.3, 1.3)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(0.5), Inches(1.3), Inches(4.0), Inches(4.8), cd)
    ch = gf.chart
    chart_common(ch)
    chart_title(ch, f"(a) 12개 전역 변동점, 지연 {len(xs_sim)} 포인트")
    ch.has_legend = False
    ser = ch.plots[0].series[0]
    ser.smooth = False
    ser.format.line.fill.background()
    ser.marker.style = XL_MARKER_STYLE.CIRCLE
    ser.marker.size = 3
    ser.marker.format.fill.solid()
    ser.marker.format.fill.fore_color.rgb = BLACK
    ser.marker.format.line.color.rgb = BLACK
    ser.marker.format.line.width = Pt(0.25)
    ser = ch.plots[0].series[1]
    ser.smooth = False
    ser.marker.style = XL_MARKER_STYLE.NONE
    ser.format.line.color.rgb = GRAY
    ser.format.line.width = Pt(0.75)
    ser.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    axis_style(ch.category_axis, title="직접 시뮬레이션  D/D0", fmt="0.0", minimum=0.8, maximum=1.3, unit=0.1)
    axis_style(ch.value_axis, title="모델 예측 (logquad)  D/D0", fmt="0.0", minimum=0.8, maximum=1.3, unit=0.1)
    # (b) mean error per corner
    add_column_chart(slide, 4.7, 1.3, 4.8, 5.3, labels, [("선형 (lin)", e_lin), ("로그 2차 (logquad)", e_logq)],
                     "(b) 전역 변동점별 평균 오차", "평균 |오차| [% of D]", ymax=2.5, unit=0.5, rot=-60, gap=50, cat_size=8)


# ---------------------------------------------------------------------------
# 도 11 : 복원 오차와 파라미터별 기여도 비교
# ---------------------------------------------------------------------------
def fig11(prs):
    S, Gd, V0 = sens_tables()
    err_u = defaultdict(list); err_s = defaultdict(list)
    for k, d in S.items():
        if k not in Gd:
            continue
        sagg = Gd[k]; v0 = V0[k]; D = 3 * SG[(k[4], k[5])]
        tot_abs = sum(abs(v) for v in d.values()); tot = sum(d.values())
        for inst, si in d.items():
            cu = abs(si) / tot_abs if tot_abs > 0 else 0
            cs = si / tot if tot != 0 else 0
            key = f"{k[0].replace('_X1', '')} {k[4]}/{k[5]}"
            err_u[key].append(abs(cu * sagg - si) * D / v0 * 100)
            err_s[key].append(abs(cs * sagg - si) * D / v0 * 100)
    keys = sorted(k for k in err_u if max(err_u[k]) > 0.005)
    comp = defaultdict(lambda: defaultdict(list))
    for k, d in S.items():
        cell, arc, j, kk, g, p = k
        od = "fall" if arc.endswith("_f") else "rise"
        if (g == "N" and od != "fall") or (g == "P" and od != "rise"):
            continue
        tot_abs = sum(abs(v) for v in d.values())
        for inst, si in d.items():
            comp[(cell, arc, inst)][p].append(abs(si) / tot_abs)
    labels, cv_, cu_ = [], [], []
    for (cell, arc, inst) in sorted(comp):
        v = comp[(cell, arc, inst)]
        if "vth" not in v or "u0" not in v:
            continue
        mv = sum(v["vth"]) / len(v["vth"]); mu = sum(v["u0"]) / len(v["u0"])
        if mv in (0.0, 1.0) and mu in (0.0, 1.0):
            continue
        labels.append(f"{cell.replace('_X1', '')} {arc.split('_')[0]}{arc.split('_')[1]} {inst}")
        cv_.append(mv); cu_.append(mu)
    slide = new_slide(prs, 11)
    add_column_chart(slide, 0.5, 1.3, 4.1, 5.3, keys,
                     [("부호 없는 C(i)", [max(err_u[k]) for k in keys]), ("부호 유지 C(i)", [max(err_s[k]) for k in keys])],
                     "(a) S_rec = C(i)·S_agg 의 복원 오차", "최대 복원 오차 [% of D, 3σ]", ymax=0.6, unit=0.1, gap=80)
    add_column_chart(slide, 4.8, 1.3, 4.7, 5.3, labels, [("C_vth", cv_), ("C_u0", cu_)],
                     "(b) 구동 그룹 적층 트랜지스터의 기여도", "기여도", ymax=1.2, unit=0.2, rot=-60, gap=60, cat_size=7.5)


# ---------------------------------------------------------------------------
# 도 12 : 장치 블록도
# ---------------------------------------------------------------------------
def fig12(prs):
    cv = Canvas(new_slide(prs, 12), 12, 7.6, 6.3)
    cv.rect(0.4, 0.5, 8.6, 6.9, "", lw=1.2)
    cv.text(0.6, 7.1, "장치 (700)", 9, ha="left")
    cv.rect(0.7, 0.9, 4.6, 5.9, "", lw=1.0)
    cv.text(0.9, 6.5, "프로세서 (710)", 8.5, ha="left")
    mods = ["기여도 결정부 (711)", "집합 민감도 획득부 (712)", "민감도 복원부 (713)", "변동량 획득부 (714)",
            "정합성 검증부 (715)", "변화량 산출부 (716)", "보정 · 라이브러리 생성부 (717)"]
    for n, m in enumerate(mods):
        cv.rect(1.0, 5.6 - n * 0.72, 4.0, 0.6, m, 7.5)
    cv.rect(5.9, 5.3, 2.9, 1.3, "메모리 (720)\n명령어, 작업 데이터", 8)
    cv.rect(5.9, 2.9, 2.9, 2.0, "저장부 (730)\n넷리스트, 소자 모델\nLVF 데이터, 집합 민감도 lib\n공칭 lib, 레이아웃 추출\n변동점 정의", 6.8)
    cv.rect(5.9, 1.0, 2.9, 1.3, "입출력부 (740)", 8)
    cv.arrow(5.3, 6.05, 5.9, 6.05); cv.arrow(5.9, 5.75, 5.3, 5.75)
    cv.arrow(5.3, 4.0, 5.9, 4.0); cv.arrow(5.9, 3.7, 5.3, 3.7)
    cv.arrow(5.3, 1.65, 5.9, 1.65)
    cv.rect(9.6, 5.9, 2.2, 1.1, "회로 시뮬레이터\n특성화 도구", 7.5)
    cv.rect(9.6, 3.9, 2.2, 1.1, "배치 · 배선 도구\n(레이아웃 추출)", 7.5)
    cv.rect(9.6, 1.9, 2.2, 1.1, "STA 도구", 7.5)
    cv.arrow(9.6, 6.45, 8.8, 4.6); cv.arrow(9.6, 4.45, 8.8, 4.2); cv.arrow(8.8, 1.65, 9.6, 2.3)


# ---------------------------------------------------------------------------
def patch_theme_fonts(prs, name):
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = part._element if hasattr(part, "_element") else etree.fromstring(part.blob)
    for el in root.iter(qn("a:majorFont"), qn("a:minorFont")):
        for tag in ("a:latin", "a:ea"):
            e = el.find(qn(tag))
            if e is not None:
                e.set("typeface", name)
    if not hasattr(part, "_element"):
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def main():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    for f in (fig1, fig2, fig3, fig4, fig5, fig6, fig7, fig8, fig9, fig10, fig11, fig12):
        f(prs)
    patch_theme_fonts(prs, FONT)
    prs.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
