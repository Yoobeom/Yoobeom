# ---------------------------------------------------------------------------
# File   : make_invention_steps_pptx.py
# Author : yoobeom.kim@samsung.com
# Purpose: 4-step explanation slides for the TR-contribution x USM timing
#          variation model.  Every element is a native PowerPoint shape,
#          connector, or table so the figure can be revised in PowerPoint.
#            python3 make_invention_steps_pptx.py [output.pptx]
#          Numbers in the tables come from example/out (ngspice 42 run).
# ---------------------------------------------------------------------------
import os, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.join(HERE, "invention_4steps.pptx")

FONT = "Malgun Gothic"
MONO = "Consolas"
SLIDE_W, SLIDE_H = 13.333, 7.5

BLACK = RGBColor(0x00, 0x00, 0x00)
DARK = RGBColor(0x26, 0x26, 0x26)
NAVY = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0x59, 0x59, 0x59)
LINE = RGBColor(0x80, 0x80, 0x80)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
TINT = RGBColor(0xE4, 0xEA, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


# ---------------------------------------------------------------------------
# text helpers
# ---------------------------------------------------------------------------
def set_font(run, size, bold=False, color=DARK, mono=False):
    f = run.font
    name = MONO if mono else FONT
    f.name = name
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    if latin is not None and rPr.find(qn("a:ea")) is None:
        ea = etree.Element(qn("a:ea"))
        ea.set("typeface", FONT)
        latin.addnext(ea)


def fill_tf(tf, lines, align=PP_ALIGN.LEFT, space=4):
    """lines: list of (text, {size, bold, color, mono, align, space_before})"""
    tf.clear()
    for i, item in enumerate(lines):
        text, o = item if isinstance(item, tuple) else (item, {})
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", align)
        p.space_after = Pt(o.get("space", space))
        if o.get("space_before"):
            p.space_before = Pt(o["space_before"])
        r = p.add_run()
        r.text = text
        set_font(r, o.get("size", 10.5), o.get("bold", False),
                 o.get("color", DARK), o.get("mono", False))


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    fill_tf(tf, lines, align)
    return tb


def box(slide, x, y, w, h, lines=None, fill=WHITE, line=LINE, lw=1.0,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, pad=0.10, dash=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
        if dash:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    tf.vertical_anchor = anchor
    fill_tf(tf, lines or [""], align)
    return shp


def arrow(slide, x0, y0, x1, y1, lw=1.5, color=NAVY):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    c.line.color.rgb = color
    c.line.width = Pt(lw)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    return c


def plain_table(gf):
    tbl = gf._element.graphic.graphicData.tbl
    pr = tbl.tblPr
    pr.set("firstRow", "0"); pr.set("bandRow", "0")
    sid = pr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(pr, qn("a:tableStyleId"))
    sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"


def cell_borders(cell, color="808080", w_pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag))
        ln.set("w", str(int(Pt(w_pt)))); ln.set("cap", "flat")
        ln.set("cmpd", "sng"); ln.set("algn", "ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")


def table(slide, x, y, colws, rowhs, header, rows, fs=9.5, hfs=10,
          align_first_left=True, hdr_fill=NAVY, hdr_color=WHITE):
    allrows = ([header] if header else []) + list(rows)
    gf = slide.shapes.add_table(len(allrows), len(colws), Inches(x), Inches(y),
                                Inches(sum(colws)), Inches(sum(rowhs)))
    plain_table(gf)
    t = gf.table
    for c, cw in enumerate(colws):
        t.columns[c].width = Inches(cw)
    for r, rh in enumerate(rowhs):
        t.rows[r].height = Inches(rh)
    for r, vals in enumerate(allrows):
        hdr = bool(header) and r == 0
        for c, v in enumerate(vals):
            cell = t.cell(r, c)
            cell_borders(cell)
            cell.fill.solid()
            cell.fill.fore_color.rgb = hdr_fill if hdr else (LIGHT if (r % 2 == 0 and not hdr) else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            al = PP_ALIGN.LEFT if (c == 0 and align_first_left) else PP_ALIGN.CENTER
            fill_tf(cell.text_frame,
                    [(v, {"size": hfs if hdr else fs, "bold": hdr,
                          "color": hdr_color if hdr else DARK, "align": al, "space": 0})])
    return gf


def slide_head(prs, title, subtitle=None, notes=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    textbox(s, 0.54, 0.32, 12.3, 0.5, [(title, {"size": 24, "bold": True, "color": NAVY})])
    if subtitle:
        textbox(s, 0.54, 0.82, 12.3, 0.32, [(subtitle, {"size": 11, "color": GRAY})])
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


# ---------------------------------------------------------------------------
# slide 1 : the 4-step figure
# ---------------------------------------------------------------------------
STEPS = [
    ("STEP 1", "TR별 기여도 생성",
     ["LVF 특성화의 TR별 섭동 민감도", "S1(i, q)"],
     ["C(i) = |S1(i)| / Σ_j |S1(j)|"],
     ["기여도 표면  C(i ; slew, load)", "Σ_i C(i) = 1"]),
    ("STEP 2", "TR별 값으로 환산",
     ["USM 라이브러리의 집합 변화량", "셀 전체를 함께 흔든 결과"],
     ["TR별 값 = C(i) × 집합값"],
     ["TR별 민감도 / 변동량", "TR 단위로 사용 가능"]),
    ("STEP 3", "내부 TR 변동량 확보",
     ["LLE DM 레이아웃 추출 결과", "(경계 TR + 내부 TR)"],
     ["TR 위치별 Δp(i) 산출"],
     ["인스턴스·TR별", "ΔVth(i), ΔU0(i)"]),
    ("STEP 4", "USM + LLE 통합",
     ["C(i),  S_agg(p),  Δp(i, p)"],
     ["ΔD = Σ_i Σ_p", "C(i)·S_agg(p)·Δp(i,p)"],
     ["전역·LLE·열화를", "하나의 식으로 처리"]),
]


def slide1(prs):
    s = slide_head(prs, "기여도로 집합값을 TR 단위로 환산해 LLE DM에 연결",
                   "셀 단위로만 주어지는 집합값을 기여도로 트랜지스터 단위로 환산하여, 트랜지스터별 변동량을 쓰는 LLE DM에 연결한다.",
                   "4단계 전체 흐름. Step 1~2는 모델 생성, Step 3은 변동량 공급, Step 4는 결합.")
    X0, W, GAP = 0.54, 2.75, 0.42
    YH, HH = 1.36, 0.66                     # header
    YR, HR = 2.06, 1.04                     # rows

    def cell(x, y, label, lines, fill):
        body = [(t, {"size": 10.5, "mono": mono, "bold": mono, "space": 2})
                for t in lines for mono in (label == "처리",)]
        box(s, x, y, W, HR,
            [(label, {"size": 8.5, "bold": True, "color": GRAY, "space": 3})] + body, fill=fill)

    for i, (no, name, inp, proc, outp) in enumerate(STEPS):
        x = X0 + i * (W + GAP)
        box(s, x, YH, W, HH,
            [(no, {"size": 9.5, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "space": 1}),
             (name, {"size": 13, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "space": 0})],
            fill=NAVY, line=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        cell(x, YR, "입력", inp, WHITE)
        cell(x, YR + HR, "처리", proc, LIGHT)
        cell(x, YR + 2 * HR, "출력", outp, WHITE)
        if i < 3:
            arrow(s, x + W + 0.03, YR + 1.5 * HR, x + W + GAP - 0.03, YR + 1.5 * HR)

    yb = YR + 3 * HR + 0.34
    box(s, X0, yb, 4 * W + 3 * GAP, 1.10,
        [("결합 결과      ΔD_arc = Σ_i Σ_p  C(i ; s,l) · S_agg(p ; s,l) · Δp(i, p)",
          {"size": 13.5, "bold": True, "color": NAVY, "mono": True, "space": 6}),
         ("개별 TR 특성화 시뮬레이션 추가 없음. LLE, 전역 파라미터 이동, 열화가 Δp(i, p)만 바꾸어 같은 경로로 들어온다.",
          {"size": 11, "space": 0})],
        fill=TINT, line=NAVY, lw=1.25, pad=0.16, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, X0, yb + 1.32, 12.26, 0.3,
            [("종래: 경계 TR을 트랜지스터별로 직접 특성화. 비용이 셀당 TR 수 × 파라미터 수 × 테이블 포인트 수에 비례한다.",
              {"size": 9.5, "color": GRAY})])


# ---------------------------------------------------------------------------
# slide 2 : step 1 detail
# ---------------------------------------------------------------------------
def slide2(prs):
    s = slide_head(prs, "LVF 특성화 결과에 TR별 민감도가 이미 있다",
                   "국부 변동 σ 산출에 사용한 S1(i, q)을 정규화만 하면 기여도가 된다. 추가 시뮬레이션은 없다.",
                   "Step 1 상세. 기존 LVF 용도와 재사용 경로의 대비.")
    W = 5.92
    box(s, 0.54, 1.32, W, 0.42, [("기존 용도 — 국부 변동 통계량", {"size": 12, "bold": True, "color": WHITE})],
        fill=GRAY, line=GRAY, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 0.54, 1.74, W, 1.98,
        [("TR별 ±δ 섭동으로 중앙 차분", {"size": 10.5, "space": 3}),
         ("S1(i, q) = [ D(q+δ) - D(q-δ) ] / 2δ", {"size": 11, "mono": True, "space": 6}),
         ("독립 가정 → 제곱합의 제곱근", {"size": 10.5, "space": 3}),
         ("σ_local² = Σ_i ( S1(i,q) · σ_mm(i) )²", {"size": 11, "mono": True, "space": 6}),
         ("Liberty  ocv_sigma_cell_rise / ocv_sigma_cell_fall 에 수록", {"size": 10, "color": GRAY, "space": 0})])

    box(s, 0.54 + W + 0.35, 1.32, W, 0.42, [("재사용 — TR별 기여도", {"size": 12, "bold": True, "color": WHITE})],
        fill=NAVY, line=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 0.54 + W + 0.35, 1.74, W, 1.98,
        [("동일한 S1(i, q)을 절대값 정규화", {"size": 10.5, "space": 3}),
         ("C(i ; s,l) = |S1(i)| / Σ_j |S1(j)|", {"size": 11, "mono": True, "space": 6}),
         ("부호를 유지하면 환산 오차가 더 작다", {"size": 10.5, "space": 3}),
         ("C(i) = S1(i) / Σ_j S1(j)", {"size": 11, "mono": True, "space": 6}),
         ("slew × load 격자마다 표면으로 저장", {"size": 10, "color": GRAY, "space": 0})],
        fill=TINT, line=NAVY)

    textbox(s, 0.54, 3.98, 12.26, 0.3,
            [("NAND2_X1  A→ZN 하강 아크의 기여도 (slew × load 3×3 격자 평균, ngspice 42)",
              {"size": 10.5, "bold": True})])
    table(s, 0.54, 4.34, [2.4, 2.5, 2.5, 2.43, 2.43], [0.52, 0.56, 0.56],
          ["파라미터", "M1  (N, A 스위칭)", "M2  (N, B 고정)", "M3  (P)", "M4  (P)"],
          [["C(i),  vth", "0.827", "0.165", "0.008", "0.000"],
           ["C(i),  u0", "0.594", "0.381", "0.025", "0.000"]])
    textbox(s, 0.54, 6.12, 12.26, 0.7,
            [("적층 구조에서는 스위칭 TR이 기여도를 지배한다. 파라미터별로 값이 달라, vth와 u0의 기여도가 0.23 차이난다.",
              {"size": 10.5, "space": 4}),
             ("따라서 파라미터 독립 기여도를 쓰려면 임계값 판정을 거쳐야 한다.", {"size": 10.5, "color": GRAY})])


# ---------------------------------------------------------------------------
# slide 3 : step 2 detail
# ---------------------------------------------------------------------------
def slide3(prs):
    s = slide_head(prs, "집합값은 기여도를 곱해야 TR 단위가 된다",
                   "USM 값은 셀 전체를 동시에 섭동한 결과라 TR 식별 정보가 없다. LLE DM은 TR별 Δp를 주는데 받을 상대가 없다.",
                   "Step 2 상세. contributing_devices : all 의 한계와 복원식, 정합성 조건.")
    box(s, 0.54, 1.32, 5.9, 0.42, [("USM 라이브러리 구조", {"size": 12, "bold": True, "color": WHITE})],
        fill=GRAY, line=GRAY, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 0.54, 1.74, 5.9, 2.06,
        [("sensitivity (usm_vth) {", {"size": 10.5, "mono": True, "space": 1}),
         ("  device_param         : delta_vth ;", {"size": 10.5, "mono": True, "space": 1}),
         ("  contributing_devices : all ;", {"size": 10.5, "mono": True, "bold": True, "color": NAVY, "space": 1}),
         ("  sens_cell_fall (tmpl) { values(...) }", {"size": 10.5, "mono": True, "space": 1}),
         ("}", {"size": 10.5, "mono": True, "space": 6}),
         ("contributing_devices : all → 어느 TR의 민감도인지 알 수 없다", {"size": 10, "color": GRAY, "space": 0})],
        fill=LIGHT)

    x2, w2 = 6.75, 6.05
    for i, (head, body, mono) in enumerate([
            ("문제", "USM 값 S_agg(p)는 셀 전체의 합이다. LLE DM이 만든 TR별 Δp를 곱할 TR별 상대가 없다.", None),
            ("환산", "S_rec(i, p) = C(i ; s,l) · S_agg(p ; s,l)", True),
            ("정합성", "모든 TR에 동일한 Δp 인가 시   Σ_i C(i)·S_agg(p)·Δp = S_agg(p)·Δp", True)]):
        y = 1.32 + i * 0.90
        box(s, x2, y, 0.95, 0.80, [(head, {"size": 11, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER})],
            fill=NAVY if i else GRAY, line=NAVY if i else GRAY, anchor=MSO_ANCHOR.MIDDLE)
        box(s, x2 + 0.95, y, w2 - 0.95, 0.80,
            [(body, {"size": 10.5, "mono": bool(mono), "bold": bool(mono)})],
            fill=WHITE, anchor=MSO_ANCHOR.MIDDLE)

    textbox(s, 0.54, 4.12, 12.26, 0.3,
            [("환산 정확도 — 그룹 동시 섭동으로 얻은 직접 민감도 대비 (3σ, 지연 대비 %)",
              {"size": 10.5, "bold": True})])
    table(s, 0.54, 4.48, [3.4, 2.9, 3.0, 2.96], [0.52, 0.56, 0.56],
          ["기여도 정의", "NAND2 최대 오차", "NOR2 최대 오차", "판정"],
          [["부호 유지  C(i) = S1(i)/ΣS1(j)", "0.005 %", "0.060 %", "권장"],
           ["부호 없음  C(i) = |S1(i)|/Σ|S1(j)|", "0.535 %", "0.542 %", "조건부"]])
    textbox(s, 0.54, 6.26, 12.26, 0.6,
            [("환산이 정합성 조건을 만족하므로, 모든 TR을 같은 양으로 움직이면 집합값이 그대로 복구된다.",
              {"size": 10.5, "space": 4}),
             ("1차 항 기준 그룹 동시 섭동과의 차이는 지연 대비 최대 0.07 % (3σ).", {"size": 10.5, "color": GRAY})])


# ---------------------------------------------------------------------------
# slide 4 : step 3 detail
# ---------------------------------------------------------------------------
def slide4(prs):
    s = slide_head(prs, "내부 TR까지 반영해야 오차가 닫힌다",
                   "LLE DM이 내부 TR의 Δp를 주더라도, 대응하는 TR별 민감도가 없으면 STA에 전달할 수 없다.",
                   "Step 3 상세. 경계 TR 한정 종래 방식과의 대비.")
    # cell schematic
    textbox(s, 0.54, 1.30, 4.6, 0.28, [("셀 내 트랜지스터 위치", {"size": 10.5, "bold": True})])
    ox, oy, ow, oh = 0.60, 1.66, 4.35, 2.50
    box(s, ox, oy, ow, oh, [""], fill=WHITE, line=DARK, lw=1.25)
    textbox(s, ox + 0.08, oy + 0.08, 1.6, 0.24, [("셀 경계", {"size": 9, "color": GRAY})])
    for i, (bx, kind) in enumerate([(0.70, "경계"), (1.75, "내부"), (2.75, "내부"), (4.12, "경계")]):
        shaded = kind == "경계"
        box(s, bx, oy + 0.68, 0.72, 1.10,
            [(f"TR{i+1}", {"size": 9.5, "bold": True, "align": PP_ALIGN.CENTER})],
            fill=RGBColor(0xD9, 0xD9, 0xD9) if shaded else WHITE, line=DARK,
            anchor=MSO_ANCHOR.MIDDLE, pad=0.02)
        textbox(s, bx - 0.06, oy + 1.84, 0.84, 0.24,
                [(kind, {"size": 9, "color": GRAY, "align": PP_ALIGN.CENTER})], align=PP_ALIGN.CENTER)
    textbox(s, 0.60, 4.28, 4.5, 0.5,
            [("음영 = 경계 TR. 종래 방식의 보정 대상은 여기까지다.", {"size": 10, "color": GRAY})])

    table(s, 5.35, 1.66, [2.05, 2.85, 2.85], [0.52, 0.70, 0.70, 0.86, 0.70],
          ["항목", "종래 (경계 TR 직접 특성화)", "본 발명 (기여도 × USM)"],
          [["보정 대상 TR", "경계 TR만", "경계 + 내부 전체"],
           ["민감도 획득", "TR별 직접 특성화 시뮬레이션", "기여도로 집합값을 환산"],
           ["추가 특성화 비용", "TR 수 × 파라미터 수 × 테이블 포인트 수에 비례", "없음"],
           ["내부 TR 변동", "미반영, 오차 잔존", "반영"]])
    textbox(s, 0.54, 5.20, 12.26, 0.9,
            [("내부 TR도 주변 레이아웃, 응력, 경년 열화로 파라미터가 변동한다.", {"size": 11, "space": 4}),
             ("기여도는 특성화 단계에서 셀 내 모든 TR에 대해 이미 산출되어 있으므로, 확장에 추가 비용이 들지 않는다.",
              {"size": 11, "space": 0})])


# ---------------------------------------------------------------------------
# slide 5 : step 4 detail
# ---------------------------------------------------------------------------
def slide5(prs):
    s = slide_head(prs, "변동 원인이 늘어도 특성화는 늘지 않는다",
                   "기여도와 집합 민감도는 고정하고 Δp(i, p)만 교체한다. 변동 원인마다 라이브러리를 다시 만들지 않는다.",
                   "Step 4 상세. 통합식, 변동 원인별 Δp 공급원, 검증 수치.")
    box(s, 0.54, 1.30, 12.26, 0.95,
        [("ΔD_arc = Σ_i Σ_p  C(i ; s,l) · S_agg(p ; s,l) · Δp(i, p)",
          {"size": 15, "bold": True, "color": NAVY, "mono": True, "align": PP_ALIGN.CENTER, "space": 4}),
         ("D_corrected = D_base + ΔD_arc          (인스턴스별 보정, 또는 라이브러리 테이블 갱신)",
          {"size": 11, "mono": True, "align": PP_ALIGN.CENTER, "space": 0})],
        fill=TINT, line=NAVY, lw=1.25, anchor=MSO_ANCHOR.MIDDLE)

    table(s, 0.54, 2.56, [3.3, 5.5, 1.9, 1.56], [0.48, 0.48, 0.48, 0.48, 0.48],
          ["변동 원인", "Δp(i, p) 공급원", "추가 특성화", "기여도 재사용"],
          [["국부 레이아웃 효과 (LLE)", "레이아웃 추출, 경계 + 내부 TR", "없음", "그대로"],
           ["전역 파라미터 이동", "전역 변동점 정의  ΔVth, ΔU0", "없음", "그대로"],
           ["경년 열화 (BTI)", "열화 모델  ΔVth(t)", "없음", "그대로"],
           ["응력 / WPE", "레이아웃 컨텍스트", "없음", "그대로"]])

    textbox(s, 0.54, 5.06, 12.26, 0.3,
            [("검증 — ngspice 42, 3 cell × 10 arc × 9 point", {"size": 10.5, "bold": True})])
    table(s, 0.54, 5.42, [5.6, 3.4, 3.26], [0.48, 0.48, 0.48, 0.48],
          ["검증 항목", "조건", "결과"],
          [["1차 합산 vs 그룹 동시 섭동", "3σ", "차이 최대 0.07 % of delay"],
           ["교차항 크기", "2-stack, 3σ", "0.4 % 이하"],
           ["logquad 모델 값 오차", "±3σ 복합 변동점", "평균 0.3 %, 최대 1.7 %"]])


# ---------------------------------------------------------------------------
def patch_theme_fonts(prs, name):
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = part._element if hasattr(part, "_element") else etree.fromstring(part.blob)
    for el in root.iter(qn("a:majorFont"), qn("a:minorFont")):
        for tag in ("a:latin", "a:ea"):
            e = el.find(qn(tag))
            if e is not None:
                e.set("typeface", name)
    if not hasattr(part, "_element"):
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def main():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    for f in (slide1, slide2, slide3, slide4, slide5):
        f(prs)
    patch_theme_fonts(prs, FONT)
    prs.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
